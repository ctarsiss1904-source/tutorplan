"""Rewrite only high-similarity general_tutor targets using their own F-J source."""
from __future__ import annotations
import csv, html, json, re
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).parent; GEN=ROOT/'data/generated'; REVIEW=ROOT/'review'; OUT=ROOT/'output'
def clean(x): return ' '.join(str(x or '').split())
def csvout(path,rows,fields):
 with path.open('w',newline='',encoding='utf-8-sig') as f:
  w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');w.writeheader();w.writerows(rows)
def main():
 pairs=list(csv.DictReader((REVIEW/'general_tutor_content_pairwise_similarity.csv').open(encoding='utf-8-sig')))
 high=[x for x in pairs if x['risk_level'] in ('HIGH','CRITICAL')]
 related=defaultdict(list)
 for x in high:
  related[x['page_id_a']].append((x['page_id_b'],x['normalized_similarity']))
  related[x['page_id_b']].append((x['page_id_a'],x['normalized_similarity']))
 pages={x['page_id']:x for x in json.loads((GEN/'nationwide_page_inventory_candidates.json').read_text(encoding='utf8'))}
 content={x['page_id']:x for x in json.loads((GEN/'general_tutor_content.json').read_text(encoding='utf8'))}
 targets=[]
 for pid in sorted(related):
  p=pages[pid]; targets.append({'page_id':pid,'region_id':p['region_id'],'title':p['title'],'content_source_id':p['content_source_id'],'high_similarity_with':'|'.join(x[0] for x in related[pid]),'similarity_score':'|'.join(x[1] for x in related[pid]),'rewrite_required':True})
 csvout(REVIEW/'general_tutor_similarity_rewrite_targets.csv',targets,list(targets[0]))
 wb=load_workbook(ROOT/'과외.xlsx',read_only=True,data_only=True); source={}
 for ws in wb.worksheets:
  for n,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
   if row and row[0]: source[(ws.title,n)]=tuple(clean(row[i] if len(row)>i else '') for i in range(5,10))
 wb.close()
 styles=[
 ('관찰 기록에서 출발하기','행동의 순서를 되짚기','원인을 분리해 보기','다음 회차의 확인점'),
 ('처음 멈추는 장면','그 뒤에 이어지는 선택','원인에 맞춘 조정','변화의 증거'),
 ('학습 흐름의 출발점','반복되는 장면의 의미','개입의 우선순위','다음 주의 기준'),
 ('기존 방식의 한계','새롭게 볼 행동','원인과 대응 연결','점검을 이어가는 법'),
 ('수업 전 확인할 신호','수업 중 포착할 행동','조정이 필요한 이유','수업 후 관찰'),
 ('변화가 필요한 순간','학습자가 보인 반응','원인을 좁히는 과정','작은 변화의 기록'),
 ('계획을 다시 읽는 시간','실행에서 드러난 문제','개입을 선택한 근거','다음 계획의 기준'),
 ('도움이 필요한 지점','혼자 수행하는 과정','원인에 대한 가설','독립 수행의 확인'),
 ('이전 습관의 흔적','전환을 만드는 행동','원인을 바꾸는 접근','새 흐름의 측정'),
 ('진단의 첫 단서','반복 장면의 분석','개선 방법의 적용','관찰 가능한 결과'),
 ('집에서 보이는 신호','수업에서 다룰 장면','원인의 연결 고리','공유할 변화 기준'),
 ('한 주의 시작','중간 점검','마무리 조정','다음 회차 준비'),
 ('실행 전의 상태','실행 중의 행동','방해 요인 확인','실행 후 비교'),
 ('학습 장면의 비교','달라진 선택','원인의 차이','유지할 기준'),
 ('문제의 표면','문제의 배경','개입의 순서','변화를 읽는 방법'),
 ('기록해야 할 장면','조정할 행동','원인의 재검토','다음 판단의 근거'),
 ('수행 흐름의 진단','실제 반응의 해석','개선의 적용점','점검의 마침표'),
 ('계획과 현실 사이','반복된 선택의 원인','새 기준의 실행','변화의 다음 단계')]
 for idx,t in enumerate(targets):
  p=pages[t['page_id']]; c=content[t['page_id']]; f,g,h,i,j=source[(p['content_source_sheet'],p['content_source_row'])]; hs=styles[idx]
  label=p['seo_region_label']; title=c['title']; intro=f'{title}는 {f}라는 상태를 하나의 결론으로 단정하지 않고, 실제 학습 장면에서 무엇이 먼저 나타나는지 기록하는 방식으로 접근합니다.'
  sections=[(hs[0],f),(hs[1],g),(hs[2],h),(hs[3],i),(f'{j}에 맞춘 수업 전개',j)]
  body='<p>'+html.escape(intro)+'</p>'+''.join(f'<h2>{html.escape(head)}</h2><p>{html.escape(text)}</p>' for head,text in sections)+f'<p>{html.escape(title)}에서는 위 변화를 {i}라는 기준으로 다음 회차까지 확인하며, {j}의 흐름에 맞춰 학습 계획을 조정합니다.</p>'
  path=ROOT/(next(x['output_path'] for x in json.loads((GEN/'general_tutor_deploy_inventory.json').read_text(encoding='utf8')) if x['page_id']==t['page_id']))
  doc=path.read_text(encoding='utf8')
  doc=re.sub(r'(?s)(</h1>).*?(<h2>관련 지역 과외</h2>)',lambda m:m.group(1)+body+m.group(2),doc,count=1)
  path.write_text(doc,encoding='utf8')
  c['html_content']=body; c['generated_at']='2026-09-10'; c['content_version']='general_tutor_v2_similarity_rewrite'; c['content_status']='rewritten_similarity_review'
 (GEN/'general_tutor_content.json').write_text(json.dumps(list(content.values()),ensure_ascii=False,indent=2),encoding='utf8')
 print(json.dumps({'targets':len(targets),'rewritten':len(targets)},ensure_ascii=False))
if __name__=='__main__':main()
