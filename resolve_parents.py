"""Resolve only the pre-existing unresolved Excel rows against an official source.

The official file is a verification input. It is never copied into the project
and never creates a tutorplan region by itself.
"""
from __future__ import annotations

import argparse, csv, json, re
from pathlib import Path
from openpyxl import load_workbook
import build

ALIASES = {"서울":"서울특별시", "부산":"부산광역시", "대구":"대구광역시", "인천":"인천광역시", "광주":"광주광역시", "대전":"대전광역시", "울산":"울산광역시", "세종":"세종특별자치시", "경기도":"경기도", "강원도":"강원특별자치도", "충북":"충청북도", "충남":"충청남도", "전북":"전북특별자치도", "전남":"전라남도", "경북":"경상북도", "경남":"경상남도", "제주도":"제주특별자치도"}

def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--official", required=True); args = parser.parse_args()
    wb = load_workbook(build.SOURCE, read_only=True, data_only=True)
    keywords = {(ws.title, n): str(row[0]) for ws in wb.worksheets for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2) if row[0]}
    # Rebuild the exact pre-resolution 312-row target set.  This deliberately
    # mirrors the prior extractor's source-order interpretation, plus the three
    # established Seoul sample relationships; it does not inspect newer output.
    baseline = {}
    for ws in wb.worksheets:
        current = None
        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if not row[0]: continue
            sido, sigungu, dong, level = build.infer_region(str(row[0]), ws.title)
            if level == "sigungu": current = sigungu
            elif level == "eupmyeondong" and current: sigungu = current
            if (ws.title, row_no) in {("서울", 28), ("서울", 29), ("서울", 30)}: sigungu = "강남구"
            values = [sido] + ([sigungu] if sigungu else []) + ([dong] if dong else [])
            ids = [build.slugify(v) for v in values]
            rid = "-".join(ids)
            baseline.setdefault(rid, {"region_id": rid, "parent_region_id": "-".join(ids[:-1]) if len(ids) > 1 else None, "source_sheet": ws.title, "source_row": row_no, "eupmyeondong": dong, "display_name": values[-1]})
    baseline_ids = set(baseline)
    targets = [x for x in baseline.values() if x["parent_region_id"] and x["parent_region_id"] not in baseline_ids]
    wb.close()
    official = list(csv.DictReader(Path(args.official).open(encoding="cp949"), delimiter="\t"))
    records = []
    for target in targets:
        sheet, row = target["source_sheet"], target["source_row"]
        keyword = keywords[(sheet, row)]
        name = build.clean_region(keyword).split()[-1]
        sido = ALIASES[sheet]
        candidates = []
        for item in official:
            full = item["법정동명"]
            parts = full.split()
            if item["폐지여부"] == "존재" and len(parts) >= 3 and parts[0] == sido and parts[-1] == name:
                candidates.append((" ".join(parts[1:-1]), full))
        parents = sorted({parent for parent, _ in candidates})
        cross_sheet = sorted({item["법정동명"].split()[0] for item in official if item["폐지여부"] == "존재" and len(item["법정동명"].split()) >= 3 and item["법정동명"].split()[-1] == name and item["법정동명"].split()[0] != sido})
        status = "confirmed" if len(parents) == 1 else "needs_review" if (len(parents) > 1 or cross_sheet) else "unresolved"
        parent = parents[0] if status == "confirmed" else None
        records.append({
            "source_sheet": sheet, "source_row": row, "original_keyword": keyword,
            "normalized_region_name": name, "candidate_sido": sido,
            "candidate_sigungu": parent, "candidate_parent": parent,
            "match_method": "official_legal_dong_exact_name_within_source_sheet" if parents else ("official_candidate_conflicts_with_source_sheet" if cross_sheet else "no_official_candidate_within_source_sheet"),
            "confidence": "high" if status == "confirmed" else "low",
            "reason": "Exactly one active legal-dong parent matched the source-sheet jurisdiction." if status == "confirmed" else ("Multiple active legal-dong parents matched within the source-sheet jurisdiction." if parents else ("Official candidates exist only outside the source-sheet jurisdiction." if cross_sheet else "No active legal-dong candidate matched within the source-sheet jurisdiction.")),
            "status": status, "candidate_parents": parents,
            "region_id": target["region_id"],
            "parent_region_id": (build.slugify(sheet) + "-" + build.slugify(parent)) if parent else None,
        })
    build.write_json(build.ROOT / "data" / "parent_resolution_targets.json", [{"source_sheet": r["source_sheet"], "source_row": r["source_row"]} for r in records])
    # These are pre-existing, already-approved sample relationships.  They are
    # retained verbatim and are outside this run's 312-row resolution scope.
    existing = [
        {"source_sheet": "서울", "source_row": 28, "candidate_parent": "강남구", "status": "confirmed", "parent_region_id": "seoul-gangnam"},
        {"source_sheet": "서울", "source_row": 29, "candidate_parent": "강남구", "status": "confirmed", "parent_region_id": "seoul-gangnam"},
        {"source_sheet": "서울", "source_row": 30, "candidate_parent": "강남구", "status": "confirmed", "parent_region_id": "seoul-gangnam"},
    ]
    build.write_json(build.PARENT_OVERRIDES, {"version": 1, "records": existing + records})
    build.REVIEW.mkdir(exist_ok=True)
    fields = list(records[0])
    def out(name, rows):
        with (build.REVIEW / name).open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore"); w.writeheader(); w.writerows(rows)
    out("unresolved_parent_before.csv", records)
    out("parent_resolution_proposals.csv", records)
    out("parent_resolution_confirmed.csv", [r for r in records if r["status"] == "confirmed"])
    out("parent_resolution_needs_review.csv", [r for r in records if r["status"] == "needs_review"])
    out("parent_resolution_remaining.csv", [r for r in records if r["status"] == "unresolved"])
    print(json.dumps({s: sum(r["status"] == s for r in records) for s in ("confirmed", "needs_review", "unresolved")}, ensure_ascii=False))
if __name__ == "__main__": main()
