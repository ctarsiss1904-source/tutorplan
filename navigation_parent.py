"""Navigation evidence and grouping; never rewrites public region/page identity."""
from __future__ import annotations

import csv
import html
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).parent
ISOLATED = 'candidate:tutor:kr-b300-ad6c-kr-ad11-c0b0-ad6c:region_tutor'
REVIEW_GROUP = '@review'
SEOUL_PARENT_MAP = 'data/seoul_navigation_parent_map.json'
DAEGU_PARENT_MAP = 'data/daegu_navigation_parent_map.json'
RELATED = re.compile(r'(<h2>관련 지역 과외</h2>\s*<ul>)(.*?)(</ul>)', re.S)


def csv_rows(path):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        return list(csv.DictReader(stream))


def norm(value):
    return ' '.join(str(value or '').split())


class Navigation:
    def __init__(self, root=ROOT):
        self.root = Path(root)
        self.out = self.root / 'output'
        self.regions = {r['region_id']: r for r in self.read('data/generated/regions.json')}
        self.seoul_parent_map = self.read_seoul_parent_map()
        self.daegu_parent_map = self.read_daegu_parent_map()
        inventory = self.read('data/generated/nationwide_page_inventory_candidates.json')
        scope = csv_rows(self.root / 'review/general_tutor_k_proposed_final_scope_1945.csv')
        public_ids = {s['page_id'] for s in scope}
        allowed = {s['page_id']: s for s in scope if s['generation_allowed'].lower() == 'true'}
        self.pages = {p['page_id']: p for p in inventory if p['page_id'] in allowed}
        if set(self.pages) != set(allowed):
            raise ValueError('PUBLIC_SCOPE/inventory identity mismatch')
        if any(s.get('scope_status') not in {'READY_AS_IS','READY_REMOVE_DUPLICATE_H1'} for s in allowed.values()):
            raise ValueError('GENERATION_ALLOWED contains non-ready scope status')
        self.by_region = {p['region_id']: p for p in self.pages.values()}
        self.sidos = {r['sido']: r for r in self.regions.values() if r['region_level'] == 'sido'}
        self.sigungu = defaultdict(list)
        for r in self.regions.values():
            if r['region_level'] == 'sigungu' and r['region_id'] != self.pages[ISOLATED]['region_id']:
                self.sigungu[r['sido']].append(r)
        workbook = load_workbook(self.root / '과외.xlsx', read_only=True, data_only=True)
        try:
            source = {(ws.title, n): row for ws in workbook.worksheets
                      for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2)}
        finally:
            workbook.close()
        self.source = source
        candidates = [p for p in inventory if p['page_type'] == 'region_tutor']
        k_page_ids = {p['page_id'] for p in candidates if self.has_k(source.get((p['content_source_sheet'], int(p['content_source_row'])), ()))}
        self.k_present = len(k_page_ids)
        if len(candidates) != 3381 or k_page_ids != public_ids:
            raise ValueError('CURRENT_EXCEL_K_INVENTORY/PUBLIC_SCOPE identity mismatch')
        evidence = defaultdict(list)
        reviews = defaultdict(list)
        for filename in ('parent_resolution_confirmed.csv', 'parent_second_pass_confirmed.csv'):
            for row in csv_rows(self.root / 'review' / filename):
                if row['status'] == 'confirmed':
                    key = (row['source_sheet'], int(row['source_row']))
                    reviews[key].append(row)
                    evidence[key].append(('TIER1', row, filename))
        for row in self.read('data/region_parent_overrides.json')['records']:
            if row.get('status') == 'confirmed':
                key = (row['source_sheet'], int(row['source_row']))
                evidence[key].append(('TIER2', row, 'region_parent_overrides.json'))
        self.records = {}
        self.groups = defaultdict(list)
        self.unresolved = defaultdict(list)
        self.conflicts = []
        for pid, page in self.pages.items():
            key = (page['content_source_sheet'], int(page['content_source_row']))
            s = allowed[pid]
            if key != (s['source_sheet'], int(s['source_row'])) or not self.has_k(source.get(key, ())):
                raise ValueError('Source/scope mismatch: ' + pid)
            keyword = norm(source[key][0])
            region = self.regions[page['region_id']]
            record = {'status': 'UNRESOLVED', 'parent': None, 'group': None, 'target_kind': None, 'evidence': []}
            if pid == ISOLATED:
                if key != ('대구', 222):
                    raise ValueError('Isolation identity/source mismatch')
                record['status'] = 'SOURCE_IDENTITY_ERROR'
            elif region['region_id'] in self.seoul_parent_map:
                parent = self.seoul_parent_map[region['region_id']]
                record.update(status='SEOUL_OFFICIAL_PARENT', parent=parent, group=parent,
                              target_kind='CONTENT_PAGE' if parent in self.by_region else 'STRUCTURAL_GROUP')
                record['evidence'].append(('SEOUL_OFFICIAL_PARENT', SEOUL_PARENT_MAP, parent))
            elif region['region_id'] in self.daegu_parent_map:
                parent = self.daegu_parent_map[region['region_id']]
                record.update(status='DAEGU_OFFICIAL_PARENT', parent=parent, group=parent,
                              target_kind='CONTENT_PAGE' if parent in self.by_region else 'STRUCTURAL_GROUP')
                record['evidence'].append(('DAEGU_OFFICIAL_PARENT', DAEGU_PARENT_MAP, parent))
            elif region['region_level'] == 'sido':
                record.update(status='STRUCTURAL_GROUP', group=REVIEW_GROUP)
            elif region['region_level'] == 'sigungu' and key[0] != '대구':
                # A source-sheet index, not an independently certified geographic assertion.
                record.update(status='STRUCTURAL_GROUP', group=self.sidos[key[0]]['region_id'])
            elif region['region_level'] == 'eupmyeondong':
                parents = set()
                for tier, row, filename in evidence.get(key, []):
                    originals = [row['original_keyword']] if row.get('original_keyword') else [r['original_keyword'] for r in reviews[key] if r['candidate_parent'] == row['candidate_parent']]
                    if not originals or any(norm(v) != keyword for v in originals):
                        if s.get('scope_origin') == 'NEW_K_1051':
                            continue
                        raise ValueError('Evidence keyword mismatch: ' + pid + ' ' + filename)
                    matches = [r for r in self.sigungu[key[0]] if r['display_name'] == row['candidate_parent']]
                    if len(matches) == 1:
                        parent = matches[0]['region_id']
                        if row.get('parent_region_id') and row['parent_region_id'] != parent:
                            raise ValueError('Evidence parent ID mismatch: ' + pid)
                        parents.add(parent)
                        record['evidence'].append((tier, filename, parent))
                explicit = self.explicit_parent(keyword, key[0])
                all_parents = parents | ({explicit} if explicit else set())
                if len(all_parents) > 1:
                    self.conflicts.append(pid)
                    raise ValueError('Conflicting navigation evidence: ' + pid)
                if parents or explicit:
                    parent = next(iter(all_parents))
                    record.update(status='TIER1_TIER2_CONFIRMED' if parents else 'TIER4_EXPLICIT', parent=parent, group=parent,
                                  target_kind='CONTENT_PAGE' if parent in self.by_region else 'STRUCTURAL_GROUP')
                    if explicit:
                        record['evidence'].append(('TIER4', keyword, explicit))
            if record['status'] == 'UNRESOLVED':
                # Contaminated Daegu source is deliberately not labelled as geographic Daegu.
                hub = REVIEW_GROUP if key[0] == '대구' else self.sidos[key[0]]['region_id']
                record['group'] = hub
                self.unresolved[hub].append(pid)
            elif record['status'] != 'SOURCE_IDENTITY_ERROR':
                self.groups[record['group']].append(pid)
            self.records[pid] = record
        for group in (self.groups, self.unresolved):
            for key in group:
                group[key].sort(key=lambda pid: (self.pages[pid]['title'], pid))
        self.name_sidos = defaultdict(set)
        for r in self.regions.values():
            self.name_sidos[r['display_name']].add(r.get('sido') or '')

    def read_seoul_parent_map(self):
        """Load the content-agnostic, audited Seoul navigation hierarchy."""
        data = self.read(SEOUL_PARENT_MAP)
        entries = data.get('entries', [])
        result = {}
        for entry in entries:
            region_id = entry.get('region_id')
            parent_region_id = entry.get('parent_region_id')
            if not region_id or not parent_region_id or entry.get('evidence_status') != 'OFFICIAL_CONFIRMED':
                raise ValueError('Invalid Seoul navigation parent map entry')
            if region_id in result:
                raise ValueError('Duplicate Seoul navigation parent map region_id: ' + region_id)
            if region_id not in self.regions or parent_region_id not in self.regions:
                raise ValueError('Unknown Seoul navigation parent map region')
            result[region_id] = parent_region_id
        seoul_regions = [r for r in self.regions.values() if r.get('source_sheet') == '서울' and r['region_level'] in {'sigungu', 'eupmyeondong'}]
        if len(result) != 375 or set(result) != {r['region_id'] for r in seoul_regions}:
            raise ValueError('Seoul navigation parent map must cover exactly 25 districts and 350 dongs')
        for region in seoul_regions:
            parent = self.regions[result[region['region_id']]]
            if region['region_level'] == 'sigungu':
                if result[region['region_id']] != 'seoul':
                    raise ValueError('Seoul district must have Seoul navigation parent')
            elif parent['region_level'] != 'sigungu' or parent.get('sido') != '서울':
                raise ValueError('Seoul dong must have a Seoul district navigation parent')
        return result

    def read_daegu_parent_map(self):
        data = self.read(DAEGU_PARENT_MAP)
        entries = data.get('entries', [])
        result = {}
        for entry in entries:
            region_id = entry.get('region_id')
            parent_region_id = entry.get('parent_region_id')
            if (not region_id or not parent_region_id or
                    entry.get('evidence_status') != 'PUBLIC_READY_PARENT_CONFIRMED'):
                raise ValueError('Invalid Daegu public navigation parent map entry')
            if region_id in result:
                raise ValueError('Duplicate Daegu navigation parent map region_id: ' + region_id)
            if region_id not in self.regions or parent_region_id not in self.regions:
                raise ValueError('Unknown Daegu navigation parent map region')
            result[region_id] = parent_region_id
        if len(result) != 78:
            raise ValueError('Daegu public navigation parent map must cover 8 sigungu and 70 public leaves')
        for region_id, parent_region_id in result.items():
            region = self.regions[region_id]
            parent = self.regions[parent_region_id]
            if region['region_level'] == 'sigungu':
                if parent_region_id != 'kr-b300-ad6c':
                    raise ValueError('Daegu sigungu must have Daegu navigation parent')
            elif parent['region_level'] != 'sigungu':
                raise ValueError('Daegu leaf must have a Daegu sigungu navigation parent')
        return result

    def read(self, filename):
        return json.loads((self.root / filename).read_text(encoding='utf-8-sig'))

    @staticmethod
    def has_k(row):
        return len(row) > 10 and row[10] is not None and str(row[10]).strip() != ''

    def explicit_parent(self, keyword, sheet):
        name = re.sub(r'\s*과외$', '', keyword)
        found = set()
        for region in self.sigungu[sheet]:
            for prefix in (region['display_name'], sheet + region['display_name']):
                if name.startswith(prefix):
                    rest = name[len(prefix):].strip()
                    if rest and rest.endswith(('동', '읍', '면', '리')):
                        found.add(region['region_id'])
        return next(iter(found)) if len(found) == 1 else None

    @staticmethod
    def directory(group):
        return '/regions/' if group == REVIEW_GROUP else '/regions/' + group + '/'

    def label(self, pid):
        r = self.regions[self.pages[pid]['region_id']]
        name, sido = r['display_name'], r.get('sido') or ''
        return (sido if len(self.name_sidos[name]) > 1 and sido and sido not in name else '') + name + '과외'

    def target(self, pid):
        if pid == ISOLATED:
            raise ValueError('Isolated navigation target')
        return self.label(pid), self.pages[pid]['route']

    def related_targets(self, pid):
        record = self.records[pid]
        if record['status'] == 'SOURCE_IDENTITY_ERROR':
            return []
        region_id = self.pages[pid]['region_id']
        region = self.regions[region_id]
        if region_id == 'seoul':
            return [self.target(i) for i in self.groups['seoul']]
        if region_id == 'kr-b300-ad6c':
            return [self.target(i) for i in self.groups['kr-b300-ad6c']]
        if region_id in self.seoul_parent_map:
            if region['region_level'] == 'sigungu':
                return [self.target(i) for i in self.groups[region_id]]
            parent = record['parent']
            return [self.target(self.by_region[parent]['page_id'])] + [
                self.target(i) for i in self.groups[parent] if i != pid
            ]
        if region_id in self.daegu_parent_map:
            if region['region_level'] == 'sigungu':
                return [self.target(i) for i in self.groups[region_id]]
            parent = record['parent']
            return [self.target(self.by_region[parent]['page_id'])] + [
                self.target(i) for i in self.groups[parent] if i != pid
            ]
        targets = []
        parent = record['parent']
        if parent:
            if parent in self.by_region:
                targets.append(self.target(self.by_region[parent]['page_id']))
            else:
                targets.append((self.regions[parent]['display_name'] + ' 지역 탐색', self.directory(parent)))
        targets.extend(self.target(i) for i in self.groups.get(region_id, [])[:6])
        if parent:
            targets.extend(self.target(i) for i in [i for i in self.groups[parent] if i != pid][:4])
        unique = {}
        for label, href in targets:
            unique.setdefault(href, label)
        sampled = [(label, href) for href, label in list(unique.items())[:10]]
        hub = record['group'] if record['status'] == 'UNRESOLVED' else region_id
        href = self.directory(hub)
        if href not in unique:
            sampled.append(('소속 확인 필요 지역 전체보기' if record['status'] == 'UNRESOLVED' else '지역 전체보기', href))
        return sampled

    @staticmethod
    def links(targets):
        return ''.join('<li><a href="' + html.escape(href, quote=True) + '">' + html.escape(label) + '</a></li>' for label, href in targets)

    def directory_body(self, group):
        targets = []
        if group == REVIEW_GROUP:
            targets.extend((r['display_name'] + ' 원본 지역 탐색', self.directory(r['region_id'])) for r in sorted(self.sidos.values(), key=lambda r: r['display_name']))
        else:
            own = self.by_region.get(group)
            if own and own['page_id'] != ISOLATED:
                targets.append(self.target(own['page_id']))
            region = self.regions[group]
            if region['region_level'] == 'sido':
                targets.extend(
                    (r['display_name'] + ' 지역 탐색', self.directory(r['region_id']))
                    for r in self.sigungu[region['sido']]
                    if self.groups.get(r['region_id']) and (not r['is_source_region'] or r['region_id'] in self.seoul_parent_map or r['region_id'] in self.daegu_parent_map)
                )
            targets.extend(self.target(pid) for pid in self.groups.get(group, []) if pid != (own or {}).get('page_id'))
        targets = list({href: (label, href) for label, href in targets}.values())
        result = '<h2>지역 탐색</h2><ul>' + self.links(targets) + '</ul>'
        pending = self.unresolved.get(group, [])
        if pending:
            result += '<h2>소속 확인 필요 지역</h2><p>시군구 소속을 확인 중인 원본 항목입니다. 목록 순서는 행정구역 소속을 뜻하지 않습니다.</p>'
            # All links remain real static HTML; sections avoid an undifferentiated long list.
            for offset in range(0, len(pending), 30):
                result += '<h3>목록 ' + str(offset // 30 + 1) + '</h3><ul>' + self.links(self.target(pid) for pid in pending[offset:offset+30]) + '</ul>'
        return result

    def related_html(self, pid):
        return self.links(self.related_targets(pid))


def refresh_related(nav=None, write=True):
    """Patch navigation only, preserving every byte outside the related list."""
    nav = nav or Navigation()
    pending = []
    for pid, page in nav.pages.items():
        path = nav.out / page['route'].strip('/') / 'index.html'
        old = path.read_text(encoding='utf-8')
        if len(RELATED.findall(old)) != 1:
            raise ValueError('Expected one related section: ' + pid)
        new = RELATED.sub(lambda m: m[1] + nav.related_html(pid) + m[3], old)
        if new != old:
            pending.append((path, new))
    if not write:
        return pending
    for path, text in pending:
        path.write_text(text, encoding='utf-8')
    return len(pending)
