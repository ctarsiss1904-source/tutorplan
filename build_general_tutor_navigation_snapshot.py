"""Generate the tracked public navigation runtime contract from local source inventory."""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).parent
INVENTORY = ROOT / 'data/generated/nationwide_page_inventory_candidates.json'
SCOPE = ROOT / 'review/general_tutor_k_proposed_final_scope_1945.csv'
OUTPUT = ROOT / 'data/general_tutor_navigation_runtime.json'
WORKBOOK = ROOT / '\uacfc\uc678.xlsx'
FIELDS = ('page_id', 'page_type', 'region_id', 'route', 'title', 'content_source_sheet', 'content_source_row', 'source_keyword', 'K_present')

def read_scope():
    with SCOPE.open(encoding='utf-8-sig', newline='') as stream:
        return list(csv.DictReader(stream))

def source_k_page_ids(candidates):
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        source = {(ws.title, number): row for ws in workbook.worksheets
                  for number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2)}
    finally:
        workbook.close()
    return {row['page_id'] for row in candidates
            if (values := source.get((row['content_source_sheet'], int(row['content_source_row']))))
            and len(values) > 10 and values[10] is not None and str(values[10]).strip()}

def validate_rows(inventory, scope):
    candidates = [row for row in inventory if row.get('page_type') == 'region_tutor']
    public_scope = {row['page_id']: row for row in scope}
    allowed = {pid: row for pid, row in public_scope.items() if row.get('generation_allowed', '').lower() == 'true'}
    if len(inventory) != 264073 or len(candidates) != 3381:
        raise ValueError('Source candidate count invariant failed')
    k_ids = source_k_page_ids(candidates)
    if len(k_ids) != 1945 or k_ids != set(public_scope) or len(allowed) != 1892 or len(public_scope) != 1945 or len(public_scope) - len(allowed) != 53:
        raise ValueError('Source public/K/hold count invariant failed')
    by_id = {row.get('page_id'): row for row in candidates}
    if len(by_id) != len(candidates) or set(allowed) - set(by_id):
        raise ValueError('Public scope/inventory identity mismatch')
    rows = []
    for row in candidates:
        pid = row['page_id']
        if pid not in allowed:
            continue
        scope_row = allowed[pid]
        if any(not isinstance(row.get(field), str) or not row[field].strip()
               for field in ('page_id', 'page_type', 'region_id', 'route', 'title', 'content_source_sheet')):
            raise ValueError('Required source field missing: ' + pid)
        if not row.get('content_source_row'):
            raise ValueError('Required source row missing: ' + pid)
        projected = {'page_id': pid, 'page_type': row['page_type'], 'region_id': row['region_id'], 'route': row['route'], 'title': row['title'], 'content_source_sheet': row['content_source_sheet'], 'content_source_row': int(row['content_source_row']), 'source_keyword': scope_row['source_keyword'].strip(), 'K_present': True}
        if projected['content_source_sheet'] != scope_row['source_sheet'] or projected['content_source_row'] != int(scope_row['source_row']) or not projected['source_keyword']:
            raise ValueError('Public source mapping mismatch: ' + pid)
        rows.append(projected)
    if len(rows) != 1892 or len({r['page_id'] for r in rows}) != 1892 or len({r['route'] for r in rows}) != 1892:
        raise ValueError('Snapshot identity invariant failed')
    return rows

def build(output=OUTPUT):
    inventory = json.loads(INVENTORY.read_text(encoding='utf-8-sig'))
    rows = validate_rows(inventory, read_scope())
    payload = json.dumps(rows, ensure_ascii=False, separators=(',', ':')) + '\n'
    output.write_bytes(payload.encode('utf-8'))
    return hashlib.sha256(payload.encode('utf-8')).hexdigest(), len(rows)

if __name__ == '__main__':
    digest, count = build()
    print(f'SNAPSHOT_ROWS={count}')
    print(f'SNAPSHOT_SHA256={digest}')
