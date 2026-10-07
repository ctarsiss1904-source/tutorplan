from __future__ import annotations

from html import escape
from pathlib import Path
import re
from shutil import copy2
from urllib.parse import quote

from openpyxl import load_workbook

from qa_generator import source_record
from thumbnail_rotation import SITE_URL, image_markup, image_url, publish_images


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "과외.xlsx"
OUTPUT = ROOT / "output"
INCHEON_GUIDE_IMAGE = ROOT / "assets" / "incheon-tutor-guide.jpg"

# Each worksheet lists its city/county rows first and then the local-area rows
# grouped beneath them.  A missing parent is kept as a directory page so no
# local-area row is placed under the wrong city/county.
REGIONS = {
    "인천": {"parents": [(3, 12, 39), (4, 40, 47), (5, 48, 53), (None, 54, 59, "연수구"), (6, 60, 70), (7, 71, 79), (8, 80, 102), (9, 103, 120), (10, 121, 133), (11, 134, 140)]},
    "광주": {"parents": [(3, 8, 26), (4, 27, 41), (5, 42, 57), (6, 58, 84), (7, 85, 117)]},
    "부산": {"parents": [(3, 19, 28), (4, 29, 38), (5, 39, 42), (6, 43, 50), (7, 51, 61), (8, 62, 70), (9, 71, 76), (10, 77, 81), (11, 82, 89), (12, 90, 97), (13, 98, 110), (14, 111, 131), (15, 132, 133), (16, 134, 138), (17, 139, 146), (18, 147, 151)]},
    "울산": {"parents": [(3, 8, 25), (4, 26, 44), (5, 45, 53), (6, 54, 77), (7, 78, 89)]},
    "경기도": {"parents": [(3, 28, 28), (4, 29, 55), (5, 56, 60), (6, 61, 66), (7, 67, 72), (8, 73, 77), (9, 78, 88), (10, 89, 103), (11, 104, 124), (12, 125, 145), (13, 146, 168), (14, 169, 179), (15, 180, 191), (16, 192, 198), (17, 199, 203), (18, 204, 207), (19, 208, 215), (20, 216, 235), (21, 236, 238), (22, 239, 250), (23, 251, 254), (24, 255, 259), (25, 260, 266), (26, 267, 277), (27, 278, 289)]},
    "강원도": {"parents": [(3, 16, 39), (4, 40, 47), (5, 48, 59), (6, 60, 69), (7, 70, 78), (8, 79, 87), (9, 88, 95), (10, 96, 104), (11, 105, 115), (None, 116, 120, "화천군"), (12, 121, 125), (13, 126, 131), (14, 132, 137), (15, 138, 143)]},
}

PAGE_OFFSETS = {
    "인천": 377,
    "광주": 516,
    "부산": 632,
    "울산": 782,
    "경기도": 870,
    "강원도": 1158,
    "세종": 1301,
}


def label(keyword: str) -> str:
    return keyword if keyword.endswith("과외") else f"{keyword}과외"


def url_path(*parts: str) -> str:
    return "/" + "/".join(quote(part, safe="") for part in parts) + "/"


def stylesheet() -> str:
    return """
    <style>
      :root { --ink:#172235; --muted:#69778a; --line:#eadfce; --paper:#fbf7ed; --card:#fffdf9; --accent:#176b87; }
      * { box-sizing:border-box; } body { margin:0; color:var(--ink); background:var(--paper); font-family:Arial, 'Noto Sans KR', sans-serif; line-height:1.7; }
      .shell { width:min(100% - 40px, 1080px); margin:0 auto; } header { border-bottom:1px solid var(--line); background:rgba(255,253,249,.92); } header .shell { min-height:68px; display:flex; align-items:center; justify-content:space-between; }
      .brand { color:var(--ink); text-decoration:none; font-size:1.25rem; font-weight:900; letter-spacing:-.05em; } .home { color:var(--muted); text-decoration:none; font-size:.9rem; }
      main { padding:48px 0 76px; } .crumb { margin:0 0 18px; color:var(--muted); font-size:.9rem; } .crumb a { color:inherit; text-decoration:none; } h1 { margin:0; font-size:clamp(2rem,5vw,3.25rem); letter-spacing:-.06em; }
      .guide-image { display:block; width:min(100%,724px); height:auto; margin:24px auto 0; }
      .page-thumbnail { margin:32px auto 0; } .page-thumbnail img { display:block; width:100%; height:auto; border-radius:18px; }
      .article, .directory, .faq, .native-faq { margin-top:24px; padding:30px; border:1px solid var(--line); border-radius:18px; background:var(--card); } .article h2, .native-faq h2 { line-height:1.35; letter-spacing:-.035em; } .native-faq { border-color:#d9c7a2; background:#fffaf0; }
      .faq h2 { margin:0 0 18px; font-size:1.25rem; } .faq article + article { margin-top:22px; padding-top:22px; border-top:1px solid var(--line); } .faq h3 { margin:0 0 8px; font-size:1.05rem; line-height:1.5; } .faq p { margin:0; color:#3c4a5c; }
      .directory h2 { margin:0 0 18px; font-size:1.25rem; } .grid { display:grid; grid-template-columns:repeat(4,1fr); gap:12px; } .card { display:flex; min-height:76px; align-items:center; justify-content:center; padding:16px; border:1px solid var(--line); border-radius:14px; color:var(--ink); background:#fff; font-weight:800; text-align:center; text-decoration:none; }
      .card:hover { border-color:var(--accent); color:var(--accent); } footer { padding:22px 0; border-top:1px solid var(--line); color:var(--muted); font-size:.88rem; }
      @media (max-width:720px) { .shell { width:min(100% - 28px,1080px); } main { padding-top:32px; } .article,.directory,.faq,.native-faq { padding:22px; } .grid { grid-template-columns:repeat(2,1fr); } .card { min-height:66px; } }
    </style>
    """


def document(title: str, breadcrumb: list[tuple[str, str | None]], body: str, path: str, thumbnail_index: int) -> str:
    crumb = " / ".join(f'<a href="{href}">{escape(text)}</a>' if href else escape(text) for text, href in breadcrumb)
    canonical_url = f"{SITE_URL}{path}"
    description = f"{title} 과외 학습 정보와 지역별 안내를 확인하세요."
    return f'''<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{escape(title)} | TutorPlan</title><link rel="canonical" href="{canonical_url}"><meta name="description" content="{escape(description)}"><meta property="og:type" content="website"><meta property="og:locale" content="ko_KR"><meta property="og:title" content="{escape(title)} | TutorPlan"><meta property="og:description" content="{escape(description)}"><meta property="og:url" content="{canonical_url}"><meta property="og:image" content="{image_url(thumbnail_index)}">{stylesheet()}</head><body><header><div class="shell"><a class="brand" href="/">TutorPlan</a><a class="home" href="/">홈으로</a></div></header><main class="shell"><nav class="crumb">{crumb}</nav>{body}</main><footer><div class="shell">TutorPlan · 전국 지역 과외</div></footer></body></html>'''


def write_page(parts: list[str], html: str) -> None:
    target = OUTPUT.joinpath(*parts, "index.html")
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(html, encoding="utf-8")


def publish_static_assets() -> None:
    target = OUTPUT / "assets" / INCHEON_GUIDE_IMAGE.name
    target.parent.mkdir(parents=True, exist_ok=True)
    copy2(INCHEON_GUIDE_IMAGE, target)


def guide_image() -> str:
    return '<img class="guide-image" src="/assets/incheon-tutor-guide.jpg" alt="과외 학습 안내">'


def article(content: str) -> str:
    if not content:
        return ""
    faq_start = re.search(r"(<h2>[^<]*</h2>)(?=\s*<h3>)", content)
    if not faq_start:
        faq_start = re.search(r"<h3>자주 묻는 질문</h3>|<h3>[^<]*(?:\?|나요|인가요|되나요|있나요|할까요)[^<]*</h3>", content)
    if not faq_start:
        return f'<section class="article">{content}</section>'
    before = content[:faq_start.start()]
    faq_content = content[faq_start.start():]
    return f'<section class="article">{before}</section><section class="native-faq">{faq_content}</section>'


def build_region(sheet, name: str, config: dict, start_index: int) -> tuple[int, int]:
    rows = {row: source_record(sheet, row) for row in range(2, sheet.max_row + 1) if sheet.cell(row, 1).value}
    root_source = rows[2]
    root_label = label(name)
    parents = []
    for entry in config["parents"]:
        source_row, start, end, *synthetic = entry
        children = [rows[row] for row in range(start, end + 1)]
        source = rows[source_row] if source_row else dict(children[0], keyword=synthetic[0], content="")
        parents.append((label(source["keyword"]), source, children))
    cards = "".join(f'<a class="card" href="{url_path(root_label, parent_label)}">{escape(parent_label)}</a>' for parent_label, _, _ in parents)
    page_index = start_index
    root_body = f'<h1>{escape(root_label)}</h1>{guide_image()}{article(root_source["content"])}<section class="directory"><h2>{escape(name)} 시군구별 과외</h2><div class="grid">{cards}</div></section>{image_markup(page_index, root_label)}'
    root_path = url_path(root_label)
    write_page([root_label], document(root_label, [("홈", "/"), (root_label, None)], root_body, root_path, page_index))
    page_index += 1
    child_count = 0
    for parent_label, parent_record, children in parents:
        child_cards = "".join(f'<a class="card" href="{url_path(root_label, parent_label, label(child["keyword"]))}">{escape(label(child["keyword"]))}</a>' for child in children)
        parent_body = f'<h1>{escape(parent_label)}</h1>{guide_image()}{article(parent_record["content"])}<section class="directory"><h2>{escape(parent_label.removesuffix("과외"))} 지역별 과외</h2><div class="grid">{child_cards}</div></section>{image_markup(page_index, parent_label)}'
        parent_path = url_path(root_label, parent_label)
        write_page([root_label, parent_label], document(parent_label, [("홈", "/"), (root_label, url_path(root_label)), (parent_label, None)], parent_body, parent_path, page_index))
        page_index += 1
        for child in children:
            child_label = label(child["keyword"])
            child_body = f'<h1>{escape(child_label)}</h1>{guide_image()}{article(child["content"])}{image_markup(page_index, child_label)}'
            child_path = url_path(root_label, parent_label, child_label)
            write_page([root_label, parent_label, child_label], document(child_label, [("홈", "/"), (root_label, url_path(root_label)), (parent_label, url_path(root_label, parent_label)), (child_label, None)], child_body, child_path, page_index))
            page_index += 1
            child_count += 1
    return len(parents), child_count


def build_sejong(sheet, start_index: int) -> int:
    rows = [source_record(sheet, row) for row in range(2, sheet.max_row + 1) if sheet.cell(row, 1).value]
    root_label = "세종과외"
    cards = "".join(f'<a class="card" href="{url_path(root_label, label(row["keyword"]))}">{escape(label(row["keyword"]))}</a>' for row in rows)
    root_record = dict(rows[0], keyword=root_label)
    page_index = start_index
    root_path = url_path(root_label)
    write_page([root_label], document(root_label, [("홈", "/"), (root_label, None)], f'<h1>{root_label}</h1>{guide_image()}<section class="directory"><h2>세종 지역별 과외</h2><div class="grid">{cards}</div></section>{image_markup(page_index, root_label)}', root_path, page_index))
    page_index += 1
    for row in rows:
        child_label = label(row["keyword"])
        child_path = url_path(root_label, child_label)
        write_page([root_label, child_label], document(child_label, [("홈", "/"), (root_label, url_path(root_label)), (child_label, None)], f'<h1>{escape(child_label)}</h1>{guide_image()}{article(row["content"])}{image_markup(page_index, child_label)}', child_path, page_index))
        page_index += 1
    return len(rows)


def main() -> None:
    publish_static_assets()
    publish_images()
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    built = []
    for name, config in REGIONS.items():
        parent_count, child_count = build_region(workbook[name], name, config, PAGE_OFFSETS[name])
        built.append(f"{name}: {parent_count}개 시군구, {child_count}개 지역")
    sejong_count = build_sejong(workbook["세종"], PAGE_OFFSETS["세종"])
    workbook.close()
    print("Built additional regional hierarchies: " + "; ".join(built) + f"; 세종: {sejong_count}개 지역.")


if __name__ == "__main__":
    main()
