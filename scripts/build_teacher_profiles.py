from __future__ import annotations

from collections import defaultdict
from html import escape
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PROFILE_BOOK = Path(r"C:\Users\lovel\Downloads\선생님프로필-성별추가.xlsx")
SCHOOL_BOOK = Path(r"C:\Users\lovel\Downloads\2025년 유초중등 학교별 학년별 학생수 학급수 입학 졸업 교원 직원 면적_260206W.xlsx")
OUTPUT = ROOT / "output" / "선생님프로필"


def value(item: object) -> str:
    return str(item).strip() if item not in (None, "") else ""


def compact(items: list[str]) -> list[str]:
    return list(dict.fromkeys(item for item in items if item))


def school_sido(name: str) -> str:
    return {
        "서울특별시": "서울",
        "부산광역시": "부산",
        "대구광역시": "대구",
        "인천광역시": "인천",
        "광주광역시": "광주",
        "대전광역시": "대전",
        "울산광역시": "울산",
        "세종특별자치시": "세종",
        "경기도": "경기",
        "강원특별자치도": "강원",
        "충청북도": "충북",
        "충청남도": "충남",
        "전북특별자치도": "전북",
        "전라남도": "전남",
        "경상북도": "경북",
        "경상남도": "경남",
        "제주특별자치도": "제주",
    }.get(name, name)


def photo_position(teacher_id: str, gender: str) -> tuple[str, str]:
    numeric = int("".join(char for char in teacher_id if char.isdigit()) or 1)
    is_female = gender == "여성"
    tile = 20 + ((numeric - 1) % 20) if is_female else (numeric - 1) % 20
    col, row = tile % 8, tile // 8
    return (f"{col / 7 * 100:.2f}%", f"{row / 4 * 100:.2f}%")


def style() -> str:
    return """<style>
      :root{--navy:#142846;--blue:#3478d4;--cream:#f7f5ef;--line:#e3e8ef;--ink:#172235;--muted:#647184}*{box-sizing:border-box}body{margin:0;color:var(--ink);background:var(--cream);font-family:'Noto Sans KR',Arial,sans-serif;letter-spacing:-.025em}a{color:inherit;text-decoration:none}header{background:#fff;border-bottom:1px solid var(--line)}.top,main{width:min(100% - 40px,1040px);margin:auto}.top{min-height:70px;display:flex;align-items:center;justify-content:space-between;font-weight:800}.brand{color:var(--navy);font-size:1.12rem}.back{color:var(--muted);font-size:.92rem}main{padding:44px 0 72px}.notice{color:#6b5a32;background:#fff8df;border:1px solid #efdfaa;border-radius:12px;padding:12px 16px;margin-bottom:20px;font-size:.9rem}.hero{display:grid;grid-template-columns:280px 1fr;gap:42px;align-items:center;background:#fff;border:1px solid var(--line);border-radius:26px;padding:38px;box-shadow:0 12px 35px rgba(20,40,70,.06)}.portrait{width:100%;aspect-ratio:.95;border-radius:18px;background-image:url('/assets/teacher-profile-candidates.png');background-size:800% auto;background-position:var(--photo-x) var(--photo-y);box-shadow:0 9px 18px rgba(20,40,70,.18)}.eyebrow{color:var(--blue);font-size:.92rem;font-weight:800}h1{margin:8px 0 14px;color:var(--navy);font-size:clamp(2rem,5vw,3.25rem);line-height:1.15}.intro{margin:0;color:#475568;font-size:1.1rem;line-height:1.75}.pills{display:flex;gap:8px;flex-wrap:wrap;margin-top:22px}.pill{border-radius:999px;padding:8px 12px;color:#2f5c9d;background:#edf5ff;font-size:.9rem;font-weight:700}.grid{display:grid;grid-template-columns:1fr 1fr;gap:20px;margin-top:24px}section{background:#fff;border:1px solid var(--line);border-radius:20px;padding:28px}section.wide{grid-column:1 / -1}h2{margin:0 0 18px;color:var(--navy);font-size:1.25rem}dl{display:grid;grid-template-columns:112px 1fr;gap:13px 18px;margin:0;line-height:1.55}dt{color:var(--muted);font-weight:700}dd{margin:0;font-weight:600}ul{margin:0;padding-left:20px;line-height:1.8}li+li{margin-top:6px}.schools{display:grid;grid-template-columns:1fr 1fr;gap:14px}.school-list{padding:18px;border-radius:14px;background:#f6f9fd}.school-list h3{margin:0 0 10px;color:var(--blue);font-size:1rem}.school-list ul{margin:0}.school-note{margin:14px 0 0;color:var(--muted);font-size:.88rem;line-height:1.6}.cta{display:flex;align-items:center;justify-content:space-between;gap:20px;background:var(--navy);color:#fff}.cta h2{color:#fff;margin-bottom:5px}.cta p{margin:0;color:#dbe9ff;line-height:1.6}.cta a{white-space:nowrap;padding:13px 18px;border-radius:10px;background:#fff;color:var(--navy);font-weight:800}.source{margin:20px 0 0;color:var(--muted);font-size:.82rem;line-height:1.6}.profile-grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(230px,1fr));gap:14px}.profile-card{display:block;padding:20px;background:#fff;border:1px solid var(--line);border-radius:16px}.profile-card strong{display:block;color:var(--navy);font-size:1.1rem;margin-bottom:6px}.profile-card span{display:block;color:var(--muted);font-size:.9rem;line-height:1.55}@media(max-width:720px){.top,main{width:min(100% - 28px,1040px)}main{padding-top:28px}.hero{grid-template-columns:1fr;gap:24px;padding:24px}.portrait{width:220px}.grid,.schools{grid-template-columns:1fr}section.wide{grid-column:auto}.cta{display:block}.cta a{display:inline-block;margin-top:18px}}
    </style>"""


def page(title: str, body: str, description: str) -> str:
    return f'''<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{escape(title)} | TutorPlan</title><meta name="description" content="{escape(description)}">{style()}</head><body><header><div class="top"><a class="brand" href="/">TutorPlan</a><a class="back" href="/선생님프로필/">선생님 프로필 목록</a></div></header><main>{body}</main></body></html>'''


def load_schools() -> dict[tuple[str, str], dict[str, list[tuple[str, str]]]]:
    workbook = load_workbook(SCHOOL_BOOK, read_only=True, data_only=True)
    sheet = workbook["학교별 주요통계"]
    schools: dict[tuple[str, str], dict[str, list[tuple[str, str]]]] = defaultdict(lambda: defaultdict(list))
    for row in sheet.iter_rows(min_row=18, values_only=True):
        sido, sigungu, level, name, status, address = value(row[1]), value(row[2]), value(row[4]), value(row[8]), value(row[15]), value(row[19])
        if level not in {"중학교", "고등학교", "초등학교"} or status == "폐교(원교)" or not name:
            continue
        schools[(sido, sigungu)][level].append((name, address))
    workbook.close()
    return schools


def select_schools(candidates: list[tuple[str, str]], dongs: list[str]) -> list[str]:
    nearby = [school for school in candidates if any(dong in school[1] for dong in dongs)]
    return compact([name for name, _ in nearby + candidates])[:2]


def school_markup(locations: list[dict[str, str]], schools: dict[tuple[str, str], dict[str, list[tuple[str, str]]]]) -> str:
    selected: dict[str, list[str]] = {}
    for level in ("중학교", "고등학교"):
        choices: list[str] = []
        for location in locations:
            choices.extend(select_schools(schools[(school_sido(location["sido"]), location["sigungu"])][level], [location["dong"]]))
        selected[level] = compact(choices)[:2]
    if not selected["중학교"] and not selected["고등학교"]:
        choices: list[str] = []
        for location in locations:
            choices.extend(select_schools(schools[(school_sido(location["sido"]), location["sigungu"])]["초등학교"], [location["dong"]]))
        selected["초등학교"] = compact(choices)[:2]
    blocks = "".join(f'<div class="school-list"><h3>{escape(level)}</h3><ul>{"".join(f"<li>{escape(name)}</li>" for name in names)}</ul></div>' for level, names in selected.items() if names)
    if not blocks:
        return ""
    return f'<section class="wide"><h2>수업 가능 지역의 학교 정보</h2><div class="schools">{blocks}</div><p class="school-note">2025 학교별 주요 통계에서 수업 가능 지역과 같은 시·군·구에 등록된 학교만 예시로 표시했습니다. 학교별 수업 가능 여부나 시험 정보를 뜻하지 않습니다.</p></section>'


def main() -> None:
    schools = load_schools()
    workbook = load_workbook(PROFILE_BOOK, read_only=True, data_only=True)
    base = {value(row[0]): row for row in workbook["학부모프로필"].iter_rows(min_row=2, values_only=True) if value(row[0])}
    grades: dict[str, list[str]] = defaultdict(list)
    for row in workbook["과목학년"].iter_rows(min_row=2, values_only=True):
        if value(row[0]): grades[value(row[0])].append(value(row[3]))
    locations: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in workbook["지역별선생님"].iter_rows(min_row=2, values_only=True):
        if value(row[0]): locations[value(row[0])].append({"sido": value(row[2]), "sigungu": value(row[3]), "dong": value(row[4])})
    workbook.close()

    OUTPUT.mkdir(parents=True, exist_ok=True)
    cards: list[str] = []
    for teacher_id, row in base.items():
        _, name, education, subject, max_grade, learning_style, method, bio, gender = map(value, row[:9])
        teacher_locations = locations[teacher_id]
        regions = compact([f'{entry["sigungu"]} {entry["dong"]}' for entry in teacher_locations])
        sido_values = compact([entry["sido"] for entry in teacher_locations])
        visit_regions = compact([f'{entry["sigungu"]} {entry["dong"]}' for entry in teacher_locations if entry["sido"] in {"서울특별시", "경기도"}])
        online_regions = compact([f'{entry["sigungu"]} {entry["dong"]}' for entry in teacher_locations if entry["sido"] not in {"서울특별시", "경기도"}])
        mode = "방문·화상 병행" if visit_regions and online_regions else ("방문 수업" if visit_regions else "화상 수업")
        x, y = photo_position(teacher_id, gender)
        route = f'/선생님프로필/{teacher_id}/'
        eyebrow = f'{regions[0]} {subject or "과외"}' if regions else (subject or "선생님 프로필")
        intro = bio or method or "등록된 선생님 프로필 정보를 확인하세요."
        profile_facts = ''.join([
            f'<dt>학력·전공</dt><dd>{escape(education or "등록 정보 없음")}</dd>',
            f'<dt>지도 과목</dt><dd>{escape(subject or "등록 정보 없음")}</dd>',
            f'<dt>지도 가능 학년</dt><dd>{escape(" / ".join(compact(grades[teacher_id])) or max_grade or "등록 정보 없음")}</dd>',
            f'<dt>최고 지도 학년</dt><dd>{escape(max_grade or "등록 정보 없음")}</dd>',
        ])
        location_facts = ''.join([
            f'<dt>수업 방식</dt><dd>{mode}</dd>',
            f'<dt>시·도</dt><dd>{escape(" · ".join(sido_values) or "등록 정보 없음")}</dd>',
            f'<dt>방문 가능 지역</dt><dd>{escape(" · ".join(visit_regions) or "해당 없음")}</dd>',
            f'<dt>화상 수업 지역</dt><dd>{escape(" · ".join(online_regions) or "해당 없음")}</dd>',
        ])
        learning_section = ""
        if learning_style or method:
            items = ''.join(f'<li>{escape(text)}</li>' for text in compact([learning_style, *method.splitlines()]))
            learning_section = f'<section class="wide"><h2>수업에서 함께 확인하는 점</h2><ul>{items}</ul></section>'
        body = f'''<p class="notice">선생님 프로필 예시 화면입니다. 프로필 사진은 예시 이미지이며 실제 선생님 사진을 뜻하지 않습니다.</p><article class="hero"><div class="portrait" role="img" aria-label="선생님 프로필 예시 이미지" style="--photo-x:{x};--photo-y:{y}"></div><div><div class="eyebrow">{escape(eyebrow)}</div><h1>{escape(name or teacher_id)} 선생님</h1><p class="intro">{escape(intro)}</p><div class="pills"><span class="pill">{escape(subject or "과목 정보 없음")}</span><span class="pill">{escape(max_grade or "학년 정보 없음")}</span><span class="pill">{escape(gender or "성별 정보 없음")}</span><span class="pill">{mode}</span></div></div></article><div class="grid"><section><h2>기본 정보</h2><dl>{profile_facts}</dl></section><section><h2>수업 방식·가능 지역</h2><dl>{location_facts}</dl></section>{school_markup(teacher_locations, schools)}{learning_section}<section class="wide cta"><div><h2>수업 전 확인할 내용</h2><p>학생의 현재 학년, 학습 상황, 가능한 지역과 일정을 확인한 뒤 수업 방향을 정합니다.</p></div><a href="tel:01029421904">상담 문의 010-2942-1904</a></section></div><p class="source">선생님·과목·학년·지역 정보는 제공된 선생님 프로필 자료를, 학교 정보는 제공된 2025 학교별 주요 통계 자료를 기준으로 구성했습니다.</p>'''
        target = OUTPUT / teacher_id / "index.html"
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(page(f'{name or teacher_id} 선생님', body, f'{name or teacher_id} 선생님 프로필과 수업 가능 지역 정보를 확인하세요.'), encoding="utf-8")
        cards.append(f'<a class="profile-card" href="{route}"><strong>{escape(name or teacher_id)} 선생님</strong><span>{escape(subject or "과목 정보 없음")} · {escape(" · ".join(regions[:2]) or "지역 정보 없음")}</span></a>')

    listing = f'<h1>선생님 프로필</h1><p class="intro">등록된 선생님 프로필을 확인하세요.</p><div class="profile-grid">{"".join(cards)}</div>'
    (OUTPUT / "index.html").write_text(page("선생님 프로필", listing, "TutorPlan 선생님 프로필 목록입니다."), encoding="utf-8")
    print(f"Created {len(base)} teacher profile pages.")


if __name__ == "__main__":
    main()
