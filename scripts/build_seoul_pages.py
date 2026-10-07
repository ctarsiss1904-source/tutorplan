from __future__ import annotations

from html import escape
from pathlib import Path
import re
from shutil import copy2
from urllib.parse import quote

from openpyxl import load_workbook

from qa_generator import source_record


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "과외.xlsx"
OUTPUT = ROOT / "output"
GUIDE_IMAGE = ROOT / "assets" / "incheon-tutor-guide.jpg"

# The Seoul worksheet is grouped by district immediately after its 25 district rows.
DISTRICTS = [
    ("강남구과외", 28, 41),
    ("강동구과외", 42, 50),
    ("강북구과외", 51, 54),
    ("강서구과외", 55, 67),
    ("관악구과외", 68, 70),
    ("광진구과외", 71, 77),
    ("구로구과외", 78, 87),
    ("금천구과외", 88, 90),
    ("노원구과외", 91, 95),
    ("도봉구과외", 96, 99),
    ("동대문구과외", 100, 109),
    ("동작구과외", 110, 117),
    ("마포구과외", 118, 141),
    ("서대문구과외", 142, 160),
    ("서초구과외", 161, 170),
    ("성동구과외", 171, 183),
    ("성북구과외", 184, 196),
    ("송파구과외", 197, 210),
    ("양천구과외", 211, 213),
    ("영등포구과외", 214, 221),
    ("용산구과외", 222, 245),
    ("은평구과외", 246, 256),
    ("종로구과외", 257, 332),
    ("서울중구과외", 333, 372),
    ("중랑구과외", 373, 378),
]


def clean_label(keyword: str, parent: str | None = None) -> str:
    label = keyword.replace("서울중구", "중구")
    if parent:
        district = parent.removesuffix("과외")
        label = label.replace(f"{district} ", "", 1)
    return label


def url_path(*parts: str) -> str:
    return "/" + "/".join(quote(part, safe="") for part in parts) + "/"


def stylesheet() -> str:
    return """
    <style>
      :root { --ink:#172235; --muted:#69778a; --line:#eadfce; --paper:#fbf7ed; --card:#fffdf9; --accent:#176b87; }
      * { box-sizing:border-box; } body { margin:0; color:var(--ink); background:var(--paper); font-family:Arial, 'Noto Sans KR', sans-serif; line-height:1.7; }
      .shell { width:min(100% - 40px, 1080px); margin:0 auto; } header { border-bottom:1px solid var(--line); background:rgba(255,253,249,.92); } header .shell { min-height:68px; display:flex; align-items:center; justify-content:space-between; }
      .brand { color:var(--ink); text-decoration:none; font-size:1.25rem; font-weight:900; letter-spacing:-.05em; } .home { color:var(--muted); text-decoration:none; font-size:.9rem; }
      main { padding:48px 0 76px; } .crumb { margin:0 0 18px; color:var(--muted); font-size:.9rem; } .crumb a { color:inherit; text-decoration:none; } h1 { margin:0; font-size:clamp(2rem,5vw,3.25rem); letter-spacing:-.06em; }
      .guide-image { display:block; width:min(100%,724px); height:auto; margin:24px auto 0; }
      .intro { margin:12px 0 28px; color:var(--muted); } .article, .directory, .faq, .native-faq { margin-top:24px; padding:30px; border:1px solid var(--line); border-radius:18px; background:var(--card); } .article h2, .native-faq h2 { line-height:1.35; letter-spacing:-.035em; } .native-faq { border-color:#d9c7a2; background:#fffaf0; }
      .faq h2 { margin:0 0 18px; font-size:1.25rem; } .faq article + article { margin-top:22px; padding-top:22px; border-top:1px solid var(--line); } .faq h3 { margin:0 0 8px; font-size:1.05rem; line-height:1.5; } .faq p { margin:0; color:#3c4a5c; }
      .directory h2 { margin:0 0 18px; font-size:1.25rem; } .grid { display:grid; grid-template-columns:repeat(4,1fr); gap:12px; } .card { display:flex; min-height:76px; align-items:center; justify-content:center; padding:16px; border:1px solid var(--line); border-radius:14px; color:var(--ink); background:#fff; font-weight:800; text-align:center; text-decoration:none; }
      .card:hover { border-color:var(--accent); color:var(--accent); } footer { padding:22px 0; border-top:1px solid var(--line); color:var(--muted); font-size:.88rem; }
      @media (max-width:720px) { .shell { width:min(100% - 28px,1080px); } main { padding-top:32px; } .article,.directory,.faq,.native-faq { padding:22px; } .grid { grid-template-columns:repeat(2,1fr); } .card { min-height:66px; } }
    </style>
    """


def document(title: str, breadcrumb: list[tuple[str, str | None]], body: str) -> str:
    crumb = " / ".join(
        f'<a href="{href}">{escape(label)}</a>' if href else escape(label)
        for label, href in breadcrumb
    )
    return f"""<!doctype html>
<html lang="ko">
  <head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{escape(title)} | TutorPlan</title>{stylesheet()}</head>
  <body>
    <header><div class="shell"><a class="brand" href="/">TutorPlan</a><a class="home" href="/">홈으로</a></div></header>
    <main class="shell"><nav class="crumb">{crumb}</nav>{body}</main>
    <footer><div class="shell">TutorPlan · 서울 지역 과외</div></footer>
  </body>
</html>"""


def write_page(parts: list[str], html: str) -> None:
    target = OUTPUT.joinpath(*parts, "index.html")
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(html, encoding="utf-8")


def guide_image() -> str:
    return '<img class="guide-image" src="/assets/incheon-tutor-guide.jpg" alt="과외 학습 안내">'


def publish_static_assets() -> None:
    target = OUTPUT / "assets" / GUIDE_IMAGE.name
    target.parent.mkdir(parents=True, exist_ok=True)
    copy2(GUIDE_IMAGE, target)


def article(content: str) -> str:
    faq_start = re.search(r"(<h2>[^<]*</h2>)(?=\s*<h3>)", content)
    if not faq_start:
        faq_start = re.search(r"<h3>자주 묻는 질문</h3>|<h3>[^<]*(?:\?|나요|인가요|되나요|있나요|할까요)[^<]*</h3>", content)
    if not faq_start:
        return f'<section class="article">{content}</section>'
    before = content[:faq_start.start()]
    faq_content = content[faq_start.start():]
    return f'<section class="article">{before}</section><section class="native-faq">{faq_content}</section>'


def main() -> None:
    publish_static_assets()
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet = workbook["서울"]
    rows = {
        row: source_record(sheet, row)
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row, 1).value
    }
    workbook.close()

    seoul_keyword = rows[2]["keyword"]
    district_rows = {keyword: rows[row] for row, (keyword, _, _) in enumerate(DISTRICTS, start=3)}
    district_cards = "".join(
        f'<a class="card" href="{url_path(seoul_keyword, clean_label(keyword))}">{escape(clean_label(keyword))}</a>'
        for keyword, _, _ in DISTRICTS
    )
    seoul_body = f"""
      <h1>{escape(seoul_keyword)}</h1>
      {guide_image()}
      {article(rows[2]['content'])}
      <section class="directory"><h2>서울 구별 과외</h2><div class="grid">{district_cards}</div></section>
    """
    write_page([seoul_keyword], document(seoul_keyword, [("홈", "/"), ("서울", None), (seoul_keyword, None)], seoul_body))

    for district_keyword, start, end in DISTRICTS:
        district_label = clean_label(district_keyword)
        district_content = district_rows[district_keyword]["content"]
        children = [rows[row] for row in range(start, end + 1)]
        child_cards = "".join(
            f'<a class="card" href="{url_path(seoul_keyword, district_label, clean_label(child["keyword"], district_keyword))}">{escape(clean_label(child["keyword"], district_keyword))}</a>'
            for child in children
        )
        district_body = f"""
          <h1>{escape(district_label)}</h1>
          {guide_image()}
          {article(district_content)}
          <section class="directory"><h2>{escape(district_label.removesuffix('과외'))} 동별 과외</h2><div class="grid">{child_cards}</div></section>
        """
        district_crumb = [("홈", "/"), (seoul_keyword, url_path(seoul_keyword)), (district_label, None)]
        write_page([seoul_keyword, district_label], document(district_label, district_crumb, district_body))

        for child in children:
            child_label = clean_label(child["keyword"], district_keyword)
            child_body = f"""
              <h1>{escape(child_label)}</h1>
              {guide_image()}
              {article(child['content'])}
            """
            child_crumb = [("홈", "/"), (seoul_keyword, url_path(seoul_keyword)), (district_label, url_path(seoul_keyword, district_label)), (child_label, None)]
            write_page([seoul_keyword, district_label, child_label], document(child_label, child_crumb, child_body))

    print(f"Built Seoul hierarchy: 1 city page, {len(DISTRICTS)} district pages, {sum(end - start + 1 for _, start, end in DISTRICTS)} dong pages.")


if __name__ == "__main__":
    main()
